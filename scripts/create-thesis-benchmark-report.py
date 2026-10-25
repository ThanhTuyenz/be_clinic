"""Build measured Word tables from saved recommendation benchmark evidence."""
import argparse
import csv
import json
import math
import os
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from statistics import mean

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "reports" / ".matplotlib"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--recovery", type=Path)
    args = parser.parse_args()
    source = json.loads(args.source.read_text(encoding="utf-8"))
    recovery = json.loads(args.recovery.read_text(encoding="utf-8")) if args.recovery else None
    if not source.get("completedAt"):
        raise ValueError("Source benchmark is incomplete")
    if recovery and not recovery.get("completedAt"):
        raise ValueError("Recovery benchmark is incomplete; wait until it finishes")
    directory = args.source.parent
    rows = source["measurements"]
    quality = source["quality"]
    recovered = {row["caseId"]: row for row in recovery["quality"]} if recovery else {}
    env = source["environment"]
    is_openai = env["model"].startswith("gpt-")
    ranking_method = "rankWithOpenAI" if is_openai else "rankWithGemini"
    fmt = lambda value: "—" if value is None else f"{value:,.2f}"
    pct = lambda value: "—" if value is None else f"{value * 100:.2f}%"
    method_name = {
        "PROFILE_COLD": "Hồ sơ: cache miss",
        "BEHAVIOR_COLD": "Hành vi: cache miss",
        "PROFILE_CACHE_C1": "Cache hồ sơ, đồng thời 1",
        "PROFILE_CACHE_C5": "Cache hồ sơ, đồng thời 5",
        "PROFILE_CACHE_C10": "Cache hồ sơ, đồng thời 10",
        "NO_CONTEXT": "Chưa có ngữ cảnh",
        "FORCED_AI_UNAVAILABLE": "Chủ động tắt AI xếp hạng",
    }
    output = directory / "Bao-cao-ket-qua-hieu-nang-va-danh-gia-AI.docx"
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.1), Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.8)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        document.styles[name].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.1
    for name in ["Title", "Heading 1", "Heading 2"]:
        document.styles[name].font.color.rgb = RGBColor.from_string("006B50")

    def p(text):
        document.add_paragraph(text)

    def h(text):
        document.add_heading(text, level=1)

    def table(headers, values):
        result = document.add_table(rows=1, cols=len(headers))
        result.style = "Light Shading Accent 1"
        header = OxmlElement("w:tblHeader")
        result.rows[0]._tr.get_or_add_trPr().append(header)
        for cell, text in zip(result.rows[0].cells, headers):
            cell.text = str(text)
        for values_row in values:
            for cell, text in zip(result.add_row().cells, values_row):
                cell.text = str(text)
        for row in result.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(9)
        document.add_paragraph()

    def actual_ai(row):
        return recovered.get(row["caseId"], row)

    def macro(cases, method):
        usable = []
        for row in cases:
            candidate = actual_ai(row).get("aiMetrics") if method == "AI" else row.get("baselineMetrics") if method == "RULE" else row.get("finalMetrics")
            if candidate:
                usable.append(candidate)
        return [len(usable)] + [mean([row[key] for row in usable if row.get(key) is not None])
                                if any(row.get(key) is not None for row in usable) else None
                                for key in ["precision3", "hit3", "mrr3", "ndcg3"]]

    document.add_heading("BÁO CÁO KẾT QUẢ THỰC NGHIỆM\nHIỆU NĂNG VÀ AI GỢI Ý GÓI KHÁM", 0)
    p("Dự án VitaCare Clinic — Số liệu đo tự động trên môi trường phát triển")
    local_time = lambda value: datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone(timedelta(hours=7))).strftime("%d/%m/%Y %H:%M:%S")
    p(f"Mã lần đo: {source['runId']}. Bắt đầu: {local_time(source['startedAt'])}; hoàn tất: {local_time(source['completedAt'])} (UTC+7).")
    p("Phạm vi: gọi trực tiếp các service NestJS đang được ứng dụng sử dụng, kết nối PostgreSQL/Redis thật và LLM thật theo cấu hình ghi trong bảng môi trường. Không bao gồm HTTP, đăng nhập, trình duyệt hoặc thời gian thao tác của người dùng. Đây là số liệu benchmark kỹ thuật; không phải dữ liệu sử dụng của bệnh nhân thật.")

    h("1. Cấu hình và dữ liệu thực nghiệm")
    table(["Thuộc tính", "Giá trị ghi nhận"], [
        ["Hệ điều hành", env["platform"] + " " + env["release"]],
        ["CPU / luồng logic", f"{env['cpu']} / {env['logicalCpus']}"],
        ["RAM", f"{env['totalRamGb']} GiB"],
        ["Node.js", env["node"]],
        ["LLM", env["model"]],
        ["Ứng viên / kết quả", f"Tối đa {env.get('candidateLimit', 8)} ứng viên mỗi nguồn; {env.get('sourceResultLimit', 6)} kết quả mỗi nguồn; {env.get('finalResultLimit', 12)} gói cuối"],
        ["Phiên bản nhãn", env.get("labelVersion", "specialty-only-v1")],
        ["Embedding", env.get("embeddingModel", "Không ghi nhận trong lần đo")],
        ["Redis vector index", env.get("vectorIndex", "Không ghi nhận trong lần đo")],
        ["Timeout xếp hạng", str(env["rankingTimeoutMs"]) + " ms"],
        ["Danh mục DB", f"{env['allPackageCount']} gói tổng; {env['activeCatalogCount']} gói hoạt động và triển khai hợp lệ"],
        ["Ngữ cảnh", f"{sum(not r['caseId'].startswith('BEH-') for r in quality)} hồ sơ giả lập; {sum(r['caseId'].startswith('BEH-') for r in quality)} truy vấn hành vi giả lập"],
        ["Mã nguồn", env["commit"] + " + thay đổi working tree; công cụ và JSON đi kèm"],
    ])
    p("Hồ sơ thử nghiệm được tạo riêng, không sử dụng bệnh án của các tài khoản đang có. Hồ sơ và tài khoản benchmark được xóa sau đo; cache benchmark dùng prefix riêng. Không tạo lịch khám, thanh toán, email hoặc sự kiện tương tác để tăng CTR.")
    p("Các nhãn được kế thừa từ chuyên khoa kỳ vọng của tình huống tổng hợp và đối chiếu chính xác tên chuyên khoa đã chuẩn hóa dấu/chữ. Đây là nhãn proxy theo kịch bản, chưa được chuyên gia y tế thẩm định. Điểm số không được gọi là độ chính xác chẩn đoán hoặc độ phù hợp lâm sàng.")
    if source.get("evaluationLabels"):
        annotated = [row for row in source["evaluationLabels"] if row["source"] == "PACKAGE_CONTENT_ANNOTATION"]
        p(f"Nhãn phiên bản 2 bổ sung ánh xạ ở cấp gói cho {len(annotated)} gói chưa gắn chuyên khoa, dựa trên tên, mô tả và triệu chứng thực tế của gói. Ánh xạ được khai báo trước khi chấm và lưu cùng JSON; không sử dụng LLM hoặc kết quả xếp hạng để gán nhãn. Gói mới không có nhãn sẽ yêu cầu rà soát thay vì tự động chấm không liên quan. Các ánh xạ này vẫn là nhãn kỹ thuật, chưa được chuyên gia y tế thẩm định.")
        p("Lần đo này thay đổi cả nội dung embedding, giới hạn ứng viên/kết quả và nhãn đánh giá. Không so sánh trực tiếp điểm với báo cáo nhãn cũ để kết luận mức cải thiện do thuật toán. Muốn so sánh trước/sau phải chấm hai phiên bản trên cùng một bộ nhãn cố định.")

    h("2. Phương pháp đo hiệu năng")
    p("Độ trễ được tính từ lúc gọi getHomeRecommendations đến lúc nhận kết quả. Dùng performance.now(); p50/p95 lấy theo nearest-rank: vị trí ceil(p × n) trên dãy tăng dần. Không tính thời gian tạo/xóa fixture. Có 5 lần đọc danh mục khởi động kết nối; bài cache có thêm 5 request khởi động không tính vào mẫu.")
    p("Cache miss theo hồ sơ vẫn thực hiện truy xuất ứng viên; cache miss theo hành vi có thể không gọi AI nếu không tìm được ứng viên. Bài chủ động tắt AI chỉ tắt xếp hạng; tầng truy xuất vẫn chạy theo code thật. Các kịch bản được chạy nối tiếp trên một máy, không đại diện năng lực máy chủ production.")
    table(["Kịch bản", "n", "TB (ms)", "p50 (ms)", "p95 (ms)", "Lỗi service"], [
        [method_name.get(row["scenario"], row["scenario"]), row["count"], fmt(row["meanMs"]), fmt(row["p50Ms"]), fmt(row["p95Ms"]), row["errors"]]
        for row in source["summary"]
    ])
    table(["Kịch bản", "Gọi AI", "AI có kết quả", "Cache hit"], [
        [method_name.get(row["scenario"], row["scenario"]), row["aiCalls"], row["aiSuccess"], row["cacheHits"]]
        for row in source["summary"]
    ])
    p("Lỗi service là request ném ngoại lệ ra ngoài. AI bị quota nhưng service trả fallback vẫn là request service thành công. Hai loại tỷ lệ phải báo cáo riêng để tránh đánh đồng khả dụng hệ thống với khả dụng LLM.")
    stages = [stage for row in rows for stage in row["stages"]]
    stage_rows = []
    for name in ["embedText", "searchWithRedisSearch", ranking_method]:
        selected = [stage for stage in stages if stage["stage"] == name]
        if selected:
            stage_rows.append([name, len(selected), sum(stage["successful"] for stage in selected),
                               fmt(mean(stage["durationMs"] for stage in selected)), fmt(max(stage["durationMs"] for stage in selected))])
    table(["Tầng thực thi", "Số lần", "Có kết quả", "TB (ms)", "Lớn nhất (ms)"], stage_rows)
    vector_calls = [stage for stage in stages if stage["stage"] == "searchWithRedisSearch"]
    if vector_calls and not any(stage["successful"] for stage in vector_calls):
        p(f"Trong {len(vector_calls)} lần truy xuất RedisSearch, không lần nào trả kết quả vector cho tập ứng viên. Module đã dùng keyword matcher theo nhánh fallback. Vì vậy số liệu này KHÔNG chứng minh hiệu năng hoặc chất lượng HNSW; nó phản ánh pipeline hiện có với truy xuất fallback và xếp hạng AI khi nhà cung cấp khả dụng.")
    elif vector_calls:
        p(f"RedisSearch trả kết quả vector ở {sum(stage['successful'] for stage in vector_calls)}/{len(vector_calls)} lời gọi truy xuất. Đây là số lần tìm được ứng viên, không phải độ chính xác gợi ý; chất lượng được đánh giá riêng ở mục 4.")
    slow_embedding = [stage for stage in stages if stage["stage"] == "embedText" and stage["durationMs"] > 10000]
    if slow_embedding and not is_openai:
        p(f"Phát hiện ngoại lệ thời gian ở bước embedding hồ sơ: lớn nhất {fmt(max(stage['durationMs'] for stage in slow_embedding))} ms. Timeout xếp hạng 10 giây không bao phủ toàn pipeline và chưa áp dụng cho hàm embedText hồ sơ. Giữ nguyên ngoại lệ trong bảng p95; cần bổ sung ngân sách thời gian cho tầng này trước khi đặt SLA.")
    table(["Đồng thời", "Request", "Thời gian lô (ms)", "Throughput (request/s)"], [
        [row["concurrency"], row["count"], fmt(row["wallMs"]), fmt(row["rps"])] for row in source["workloads"]
    ])
    p("Throughput trên là thông lượng gọi service đọc cache, không phải số người dùng đồng thời hoặc số HTTP request/s mà hệ thống production hỗ trợ.")
    # Compare provider success/failure latency separately; fallback time is not LLM speed.
    latency_groups = []
    for label, subset in [
        ("Cold: AI thành công", [r for r in rows if r["scenario"] in ["PROFILE_COLD", "BEHAVIOR_COLD"] and r.get("aiSuccess", 0) > 0]),
        ("Cold: gọi AI nhưng thất bại", [r for r in rows if r["scenario"] in ["PROFILE_COLD", "BEHAVIOR_COLD"] and r.get("aiCalls", 0) > 0 and r.get("aiSuccess", 0) == 0]),
        ("Cold: không có lời gọi AI", [r for r in rows if r["scenario"] in ["PROFILE_COLD", "BEHAVIOR_COLD"] and r.get("aiCalls", 0) == 0]),
    ]:
        times = sorted(r["durationMs"] for r in subset)
        if times:
            latency_groups.append([label, len(times), fmt(mean(times)), fmt(times[math.ceil(.5 * len(times)) - 1]), fmt(times[math.ceil(.95 * len(times)) - 1])])
    table(["Nguồn thực thi cold", "n", "TB (ms)", "p50 (ms)", "p95 (ms)"], latency_groups)

    labels = ["Cache C=1", "Cache C=5", "Cache C=10", "No context"]
    plotted = [next(r for r in source["summary"] if r["scenario"] == scenario) for scenario in ["PROFILE_CACHE_C1", "PROFILE_CACHE_C5", "PROFILE_CACHE_C10", "NO_CONTEXT"]]
    fig, ax = plt.subplots(figsize=(8, 3.3))
    positions = list(range(len(labels)))
    ax.bar([x - .18 for x in positions], [r["p50Ms"] for r in plotted], .36, label="p50", color="#059669")
    ax.bar([x + .18 for x in positions], [r["p95Ms"] for r in plotted], .36, label="p95", color="#334155")
    ax.set_xticks(positions, labels)
    ax.set_ylabel("Service latency (ms)")
    ax.legend()
    fig.tight_layout()
    figure_path = directory / "latency.png"
    fig.savefig(figure_path, dpi=180)
    plt.close(fig)
    document.add_picture(str(figure_path), width=Inches(6.3))

    h("3. Độ ổn định vận hành và lỗi ghi nhận")
    cold = [row for row in rows if row["scenario"] in ["PROFILE_COLD", "BEHAVIOR_COLD"]]
    rank_calls = [stage for row in cold for stage in row["stages"] if stage["stage"] == ranking_method]
    cached = [row for row in rows if row["scenario"].startswith("PROFILE_CACHE_C")]
    forced = [row for row in rows if row["scenario"] == "FORCED_AI_UNAVAILABLE"]
    checks = [
        ("Service trả kết quả không ném lỗi", sum(row["ok"] for row in rows), len(rows)),
        ("Rerank cold trả danh sách hợp lệ", sum(stage["successful"] for stage in rank_calls), len(rank_calls)),
        ("Request bài cache dùng được cache", sum(bool(row.get("fromCache")) for row in cached), len(cached)),
        ("Chủ động tắt AI vẫn trả kết quả", sum(row["ok"] and row.get("packageCount", 0) > 0 for row in forced), len(forced)),
    ]
    table(["Chỉ tiêu", "Đạt", "Tổng quan sát", "Tỷ lệ"], [[label, passed, total, pct(passed / total) if total else "—"] for label, passed, total in checks])
    p("Độ ổn định trong bảng là khả dụng vận hành trong một lần chạy. Chưa đo mức lặp lại thứ tự gợi ý qua nhiều lần gọi cùng prompt, chưa phải kiểm thử tải dài hạn. Danh sách hợp lệ là đầu ra đã qua kiểm tra backend; không suy ra model không từng sinh nội dung bị loại.")
    if not is_openai:
        p("Lượt thăm dò trước sửa ghi nhận Gemini trả HTTP 400: deadline 5 giây quá ngắn, mức tối thiểu được phản hồi là 10 giây. Luồng xếp hạng cũ đặt 4.000 ms nên tự chuyển fallback. Timeout xếp hạng và tìm kiếm embedding có deadline ngắn đã được đổi thành 10.000 ms. Bảng chính chỉ dùng lần chạy sau sửa; thời gian từ chối trước sửa không được gọi là tốc độ AI.")
    errors = Counter((row["provider"], row["category"], str(row.get("status"))) for row in source["providerErrors"])
    table(["Nhà cung cấp / tầng", "Loại", "HTTP/mã", "Số lỗi lần đo chính"], [[*key, value] for key, value in errors.items()] or [["—", "Không ghi nhận", "—", 0]])
    p("Thực nghiệm cold lần đầu không giãn nhịp gọi. Quota là đặc tính môi trường/tài khoản sử dụng trong lần đo, không phải bằng chứng model suy luận sai. Không gộp độ trễ trả lỗi quota vào nhóm AI thành công.")
    if recovery:
        recovered_count = sum(row["aiSuccessful"] for row in recovery["quality"])
        p(f"Đã chạy bổ sung {len(recovery['measurements'])} lời gọi xếp hạng, khôi phục thành công {recovered_count}/{len(recovery['quality'])} tình huống thiếu kết quả AI. Giãn nhịp tối thiểu 15 giây giữa các lần bắt đầu, chờ thêm theo phản hồi quota khi cần; tối đa 3 lần/ca. Giữ nguyên prompt và tập ứng viên đã chụp từ lần chính.")
        p("Lần bổ sung chỉ đo tầng xếp hạng, không chạy lại truy xuất hoặc API. Kết quả được dùng để hoàn thiện bảng chất lượng trên cùng tập ứng viên; không thay thế độ trễ và tỷ lệ lỗi của lần đo chính.")
        if recovery.get("stopReason"):
            p("Giới hạn lần bổ sung: " + recovery["stopReason"])
        recovery_errors = Counter((row["provider"], row["category"], str(row.get("status"))) for row in recovery["providerErrors"])
        table(["Tầng phục hồi", "Loại lỗi", "HTTP/mã", "Số lỗi"], [[*key, value] for key, value in recovery_errors.items()] or [["—", "Không ghi nhận", "—", 0]])

    h("4. Phương pháp và kết quả đánh giá chất lượng")
    p(f"Baseline là rankWithRules của dự án; phương pháp AI là {ranking_method}. Hai phương pháp nhận cùng tập ứng viên và cùng nhãn của từng ca, không thay đổi tập ứng viên để làm tăng điểm AI. Gói liên quan được xác định bằng nhãn chuyên khoa và các ánh xạ cấp gói đã khai báo nếu có; không sử dụng LLM tự chấm chính nó.")
    p("Precision@3 = số gói liên quan trong 3 vị trí đầu / 3, vị trí thiếu tính không liên quan. Hit@3 bằng 1 nếu top 3 có gói liên quan. MRR@3 bằng nghịch đảo vị trí liên quan đầu tiên, bằng 0 nếu không có. NDCG@3 dùng nhãn nhị phân và DCG = tổng rel_i/log2(i+1); IDCG tính từ toàn bộ gói liên quan hợp lệ trong danh mục. Lấy trung bình theo ca, không gộp tất cả gói thành một mẫu độc lập.")
    p("Ba ca TC-28 đến TC-30 có ý định tái khám nhưng lần đo không dựng lịch sử khám đã hoàn tất. Chúng được báo cáo riêng; không dùng để kết luận luật tái khám đúng. Chỉ số chính bên dưới loại ba ca này để tránh đánh đồng ý định tái khám với trạng thái lịch sử đã tồn tại.")
    primary = [row for row in quality if not row.get("followUpHistoryNotSimulated") and row["caseId"] != "BEH-TC-28"]
    # Compare paired subsets, so AI/baseline denominators cannot silently differ.
    groups = []
    for label, subset in [
        ("Hồ sơ", [row for row in primary if not row["caseId"].startswith("BEH-")]),
        ("Hành vi", [row for row in primary if row["caseId"].startswith("BEH-")]),
    ]:
        paired = [row for row in subset if actual_ai(row).get("aiMetrics") and row.get("baselineMetrics")]
        groups.append((label, subset, paired))
    coverage = []
    scores = []
    for label, subset, paired in groups:
        coverage.append([label, len(subset), sum(bool(row["candidateIds"]) for row in subset), len(paired), sum(not row["candidateIds"] for row in subset)])
        for method in ["RULE", "AI"]:
            n, precision, hit, mrr, ndcg = macro(paired, method)
            scores.append([label, "Quy tắc" if method == "RULE" else "LLM", n, pct(precision), pct(hit), fmt(mrr), fmt(ndcg)])
    table(["Nguồn", "Ca chính", "Có ứng viên", "Cặp AI/baseline hợp lệ", "Không có ứng viên"], coverage)
    retrieval_coverage = []
    for label, subset, _ in groups:
        reachable = sum(bool(set(row["candidateIds"]) & set(row["relevantIds"])) for row in subset)
        retrieval_coverage.append([label, len(subset), reachable, len(subset) - reachable, pct(reachable / len(subset)) if subset else "—"])
    table(["Nguồn", "Ca chính", "Có ứng viên liên quan", "Thiếu ứng viên liên quan", "Bao phủ liên quan"], retrieval_coverage)
    p("Ứng viên liên quan được đối chiếu với cùng nhãn chuyên khoa proxy. Nếu tầng truy xuất không lấy được gói liên quan vào danh sách ứng viên, LLM không thể bổ sung gói đó vì chỉ được xếp hạng ID đã cung cấp. Đây là hạn chế của tầng truy xuất cần xem xét riêng, không được quy toàn bộ điểm thấp cho rerank.")
    table(["Nguồn", "Phương pháp", "n cặp", "P@3", "Hit@3", "MRR@3", "NDCG@3"], scores)
    p("Các ca không có ứng viên không xuất hiện trong bảng so sánh xếp hạng có điều kiện. Chúng vẫn xuất hiện trong bảng bao phủ và bảng pipeline cuối. Vì vậy, điểm xếp hạng cao không đồng nghĩa toàn bộ pipeline tìm được gói phù hợp ở mọi ca.")
    final_rows = []
    for label, subset, _ in groups:
        n, precision, hit, mrr, ndcg = macro(subset, "FINAL")
        final_rows.append([label, n, pct(precision), pct(hit), fmt(mrr), fmt(ndcg)])
    table(["Pipeline cuối lần chính", "n", "P@3", "Hit@3", "MRR@3", "NDCG@3"], final_rows)
    p("Bảng pipeline cuối dùng danh sách thật trả về trong lần chính, gồm AI/fallback/gói bổ sung theo trạng thái lúc đó; không thay bằng kết quả xếp hạng phục hồi. Hai bảng trả lời hai câu hỏi khác nhau: chất lượng xếp hạng trên ứng viên cố định và chất lượng kết quả hệ thống thực sự trả về.")
    constrained = [row for row in quality if "finalConstraintViolations" in row]
    checked_packages = sum(len(row["finalPackages"]) for row in constrained)
    violations = sum(row["finalConstraintViolations"] for row in constrained)
    table(["Kiểm tra", "Số quan sát", "Vi phạm"], [["Gói cuối theo hồ sơ, kiểm tra lại bằng bộ lọc backend", checked_packages, violations]])
    p("Đây là kiểm tra tuân thủ cùng bộ lọc backend, không phải đánh giá độc lập độ an toàn y khoa. Bộ chặn ID ngoài ứng viên có thể loại đầu ra sai; không suy ra model chưa từng sinh ID sai chỉ từ danh sách đã lọc.")

    h("5. Kết quả chi tiết từng tình huống")
    detail = []
    csv_rows = []
    for row in quality:
        ai_row = actual_ai(row)
        metrics = ai_row.get("aiMetrics") or {}
        baseline = row.get("baselineMetrics") or {}
        status = "Phục hồi" if row["caseId"] in recovered and ai_row.get("aiSuccessful") else "Lần chính" if row["aiSuccessful"] else "Không có AI"
        detail.append([row["caseId"], len(row["candidateIds"]), status, pct(baseline.get("precision3")), pct(metrics.get("precision3")), fmt(metrics.get("ndcg3"))])
        csv_rows.append({"caseId": row["caseId"], "candidateCount": len(row["candidateIds"]), "aiSource": status,
                         "relevantCount": len(row["relevantIds"]), "baselineP3": baseline.get("precision3"), "aiP3": metrics.get("precision3"),
                         "aiHit3": metrics.get("hit3"), "aiMrr3": metrics.get("mrr3"), "aiNdcg3": metrics.get("ndcg3"),
                         "expectedSpecialties": "; ".join(row["expectedSpecialties"])})
    table(["Ca", "Ứng viên", "Nguồn AI", "P@3 quy tắc", "P@3 AI", "NDCG AI"], detail)

    h("6. Nhận xét và giới hạn")
    first = source["summary"][0]
    cache = next(row for row in source["summary"] if row["scenario"] == "PROFILE_CACHE_C1")
    p(f"Trong lần đo, p50 gọi service đọc cache với đồng thời 1 là {fmt(cache['p50Ms'])} ms; p95 là {fmt(cache['p95Ms'])} ms trên {cache['count']} request. Đây là kết quả đo trên máy được mô tả ở bảng cấu hình, không phải cam kết SLA.")
    for label, subset, paired in groups:
        if paired:
            baseline = macro(paired, "RULE")
            llm = macro(paired, "AI")
            p(f"Nhóm {label.lower()} có {len(paired)} cặp hợp lệ: Precision@3 quy tắc {pct(baseline[1])}, AI {pct(llm[1])}; chênh lệch {(llm[1] - baseline[1]) * 100:+.2f} điểm phần trăm. Đây là so sánh mô tả trên nhãn proxy; chưa kiểm định ý nghĩa thống kê hoặc xác nhận phù hợp y khoa.")
    for text in [
        "Không có người dùng thật trong benchmark này; không tính hoặc suy diễn CTR, mức hài lòng, tỷ lệ mua/đặt tự nhiên.",
        "Một lần chạy với số ca nhỏ, không có lặp nhiều lần cho từng ca hoặc khoảng tin cậy; p95 nhóm AI nhỏ cần diễn giải thận trọng.",
        "Các gói cùng chuyên khoa ở nhiều chi nhánh có thể cùng được chấm liên quan. Điểm không phản ánh độ đa dạng hoặc lựa chọn chi nhánh của người dùng.",
        "Các ca dùng nhãn chuyên khoa soạn sẵn, chưa tách một tập test độc lập với quá trình phát triển; chưa có nhãn thẩm định chuyên gia.",
        "Không loại ca quota khỏi số liệu độ tin cậy lần chính; chỉ dùng cặp có kết quả AI khi so sánh xếp hạng. Báo cáo bao phủ mẫu riêng để không che ca thất bại.",
    ]:
        p(text)

    h("7. Bằng chứng và tái lập")
    p("File dữ liệu chính: " + str(args.source.resolve()))
    if args.recovery:
        p("File xếp hạng phục hồi: " + str(args.recovery.resolve()))
    p("Công cụ: scripts/benchmark-recommendations.cjs; scripts/create-thesis-benchmark-report.py. Mỗi request lưu thời gian, cache hit, lời gọi các tầng và nguồn thực thi. Phần chất lượng lưu tập ứng viên, nhãn và thứ tự kết quả để kiểm tra lại.")
    p("Lệnh đo: node scripts/benchmark-recommendations.cjs. Lệnh phục hồi có giãn nhịp: node scripts/benchmark-recommendations.cjs --recover <raw-results.json>. Mỗi lần chạy có thể dùng quota API; kết quả không tất định và có thể khác khi danh mục/model/mạng thay đổi.")
    p("Số liệu trong các bảng được sinh từ JSON đã lưu, không điền thủ công hoặc thay số khi kết quả không thuận lợi.")
    document.save(output)
    checked = Document(output)
    assert len(checked.tables) >= 10 and len(checked.inline_shapes) == 1
    with (directory / "performance.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(source["summary"][0]))
        writer.writeheader()
        writer.writerows(source["summary"])
    with (directory / "quality.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(csv_rows[0]))
        writer.writeheader()
        writer.writerows(csv_rows)
    summary = {"word": str(output), "requests": len(rows), "aiInitialSuccess": sum(row.get("aiSuccess", 0) for row in rows),
               "aiInitialCalls": sum(row.get("aiCalls", 0) for row in rows),
               "aiAvailableQualityCases": sum(actual_ai(row).get("aiSuccessful", False) for row in quality),
               "qualityCases": len(quality), "coverage": coverage, "retrievalCoverage": retrieval_coverage,
               "pairedQuality": scores, "tables": len(checked.tables)}
    (directory / "report-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
