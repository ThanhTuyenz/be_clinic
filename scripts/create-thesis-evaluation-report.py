"""Create the Word evaluation plan; do not invent benchmark results."""
from datetime import datetime
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "thesis" / "Ke-hoach-do-hieu-nang-va-danh-gia-AI.docx"


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ["Title", "Heading 1", "Heading 2"]:
        document.styles[name].font.name = "Times New Roman"
        document.styles[name].font.color.rgb = RGBColor.from_string("006B50")

    def paragraph(text):
        document.add_paragraph(text)

    def heading(text):
        document.add_heading(text, level=1)

    def table(headers, rows):
        result = document.add_table(rows=1, cols=len(headers))
        result.style = "Light Shading Accent 1"
        for cell, value in zip(result.rows[0].cells, headers):
            cell.text = value
        for row in rows:
            for cell, value in zip(result.add_row().cells, row):
                cell.text = str(value)
        for row in result.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)
        document.add_paragraph()

    document.add_heading("KẾ HOẠCH ĐO HIỆU NĂNG\nVÀ ĐÁNH GIÁ AI GỢI Ý GÓI KHÁM", 0)
    paragraph("Dự án VitaCare Clinic — Tài liệu chuẩn bị thực nghiệm cho khóa luận")
    paragraph("Ngày lập: " + datetime.now().strftime("%d/%m/%Y %H:%M"))
    paragraph("Trạng thái: đã rà soát mã nguồn; chưa chạy phép đo hiệu năng và chưa xác lập bộ nhãn chất lượng AI. Tài liệu này là kế hoạch, không phải báo cáo kết quả thực nghiệm.")

    heading("1. Mục tiêu và phạm vi")
    paragraph("Thực nghiệm có hai mục tiêu độc lập: (1) đo thời gian phản hồi và khả năng xử lý của module gợi ý; (2) đánh giá chất lượng danh sách gói do AI xếp hạng so với phương pháp đối chứng. Các chỉ số CTR và chuyển đổi từ người dùng không thay thế hai mục tiêu này.")
    paragraph("Kết quả kỹ thuật có thể thu thập tự động, không cần tuyển người tham gia. Đánh giá mức độ liên quan cần bộ nhãn rõ nguồn gốc; nhãn kỹ thuật theo chuyên khoa hoặc từ khóa chỉ là đại diện gần đúng, chưa phải xác nhận phù hợp y khoa.")

    heading("2. Kiến trúc được xác minh từ mã nguồn")
    table(["Thành phần", "Hành vi hiện tại"], [
        ["Nguồn hồ sơ", "Lọc ứng viên theo hồ sơ; phân tích và cache độc lập theo hồ sơ."],
        ["Nguồn hành vi", "Phân tích theo tài khoản/phiên, tách khỏi hồ sơ đang chọn."],
        ["LLM", "Xếp hạng ứng viên đã truy xuất; backend kiểm tra ID trong danh sách đầu vào."],
        ["Kết hợp", "Backend gộp hai danh sách, khử trùng, kiểm tra lại điều kiện và bổ sung gói hợp lệ."],
        ["Giới hạn", "Mỗi nguồn tối đa 6 gói; danh sách cuối tối đa 12 gói."],
        ["Cache", "Hồ sơ 7 ngày; hành vi 24 giờ; fallback theo quy tắc tối đa 5 phút."],
        ["Mô hình và timeout", "Mặc định gemini-2.5-flash; timeout xếp hạng 10.000 ms sau khi sửa lỗi deadline. Cấu hình thực tế phải ghi lại khi đo."],
    ])
    paragraph("Không chỉ có trường hợp đủ điểm chuyên khoa: mã hiện tại cũng kích hoạt nhánh hành vi khi có truy vấn tìm kiếm hoặc gói bỏ dở hợp lệ. Vì vậy, mỗi ca đo cần ghi rõ ngữ cảnh kích hoạt.")

    heading("3. Rà soát công cụ đánh giá có sẵn")
    source = (ROOT / "scripts" / "evaluate-accuracy.mjs").read_text(encoding="utf-8")
    count = len(re.findall(r"id:\s*'TC-\d+'", source))
    paragraph(f"Đã tìm thấy scripts/evaluate-accuracy.mjs với {count} tình huống định nghĩa trong mã. Script tự dựng các bước truy xuất và gọi AI, thay vì gọi trực tiếp toàn bộ service đang được ứng dụng sử dụng. Không chạy script này để mặc nhiên đại diện cho luồng hiện tại.")
    paragraph("Bộ chấm đang dựa trên từ khóa/chuyên khoa và từ khóa cấm. Chưa có bằng chứng bộ nhãn đã được chuyên gia thẩm định; không gọi bộ này là chuẩn lâm sàng. Một số kiểm tra tên gói, tái khám và cách tính Precision@3 cần đối chiếu lại trước khi tái sử dụng.")

    heading("4. Thiết kế đo hiệu năng")
    paragraph("Đo riêng độ trễ toàn API và độ trễ các bước quan trọng: đọc dữ liệu, truy xuất ứng viên, gọi embedding nếu có, gọi LLM, đọc/ghi cache và gộp kết quả. Không so sánh thời gian hàm nội bộ với thời gian HTTP như cùng một đại lượng.")
    table(["Mã", "Kịch bản", "Điều kiện phải xác nhận"], [
        ["P01", "Gợi ý hồ sơ chưa có cache", "Cache miss; ghi nguồn AI thực tế hoặc fallback."],
        ["P02", "Gợi ý hồ sơ đã có cache", "Cùng ngữ cảnh; cache hit; không gọi lại LLM."],
        ["P03", "Gợi ý hành vi chưa có cache", "Hành vi đủ điều kiện kích hoạt; không kèm hồ sơ trong prompt."],
        ["P04", "Gợi ý hành vi đã có cache", "Cùng hành vi và danh mục; cache hit."],
        ["P05", "Hồ sơ và hành vi cùng hoạt động", "Ghi cache hit/miss và nguồn thực thi của từng nhánh."],
        ["P06", "Không có ngữ cảnh cá nhân hóa", "Danh sách bổ sung; không gán kết quả này cho LLM."],
        ["P07", "LLM lỗi hoặc quá thời gian", "Lỗi được chủ động mô phỏng; ghi riêng kết quả fallback."],
    ])
    paragraph("Quy mô chạy thử đề xuất, chưa phải số liệu đã thu thập: với đường cache, khởi động bằng 10 request không tính vào mẫu, sau đó đo 100 request ở từng mức đồng thời 1, 5 và 10. Với đường gọi LLM mới, thăm dò bằng 5 request; sau khi kiểm tra quota và lỗi, nhắm 20 lần hợp lệ cho mỗi nhóm nguồn. Với mẫu nhỏ, p95 chỉ dùng mô tả và phải ghi số mẫu.")
    paragraph("Không tăng tải nhà cung cấp LLM ngay từ đầu. Tách bài chịu tải đọc cache khỏi bài đo độ trễ nhà cung cấp. Ghi rõ lần chạy đầu, khởi động mô hình và cache truy vấn embedding. Nếu chỉ đo service nội bộ, báo cáo phải ghi service-level, không gọi là HTTP API benchmark.")
    table(["Chỉ số", "Cách trình bày"], [
        ["Số request", "Tổng gửi, thành công, lỗi; không loại lỗi khỏi mẫu mà không công bố."],
        ["Độ trễ", "Trung vị (p50), p95, trung bình; đơn vị ms; ghi quy tắc tính percentile."],
        ["Thông lượng", "Số request hoàn tất / tổng thời gian đo, đơn vị request/s."],
        ["Tỷ lệ lỗi", "Số request thất bại / tổng request × 100%."],
        ["Nguồn thực thi", "AI thành công, cache hit và fallback báo cáo riêng; HTTP thành công chưa chắc AI thành công."],
    ])

    heading("5. Thiết kế đánh giá chất lượng AI")
    paragraph("Bộ ban đầu dự kiến 30 tình huống: 10 theo hồ sơ, 10 theo hành vi và 10 kiểm tra kết hợp/ràng buộc. Đây là thiết kế đề xuất, chưa phải bộ đã gán nhãn hoặc đã chạy. Đầu vào hồ sơ dùng dữ liệu giả lập được công bố; danh mục gói lấy từ database tại thời điểm thực nghiệm.")
    paragraph("Mỗi tình huống phải có: mã ca; nguồn đầu vào; dữ liệu giả lập; danh mục và ứng viên hợp lệ; nhãn mức liên quan theo packageId; nguồn/ người gán nhãn; danh sách trả về; nguồn thực thi; thời gian; lỗi. Không tạo gói hoặc mã gói giả để bù cho danh mục thiếu.")
    paragraph("Khi đo riêng tác động xếp hạng, AI và baseline theo quy tắc phải nhận cùng tập ứng viên. Khi đo toàn bộ pipeline, báo cáo thêm khả năng truy xuất được gói liên quan trước xếp hạng, để phân biệt lỗi truy xuất với lỗi LLM.")
    table(["Chỉ số", "Định nghĩa và điều kiện"], [
        ["Precision@3", "Số gói liên quan trong 3 vị trí đầu / 3. Nếu thiếu kết quả, vị trí trống không tính liên quan. Báo riêng trường hợp danh mục không có ứng viên phù hợp."],
        ["Hit@3", "Tỷ lệ tình huống có ít nhất một gói liên quan trong 3 vị trí đầu."],
        ["MRR@3", "Trung bình nghịch đảo vị trí gói liên quan đầu tiên trong top 3; bằng 0 nếu không có."],
        ["NDCG@3", "So sánh thứ tự trả về với thứ tự lý tưởng; cần nhãn liên quan nhất quán. Công bố thang nhãn và xử lý ca không có gói liên quan."],
        ["Vi phạm ràng buộc", "Số gói vi phạm điều kiện cấu trúc tuổi, giới tính, hoạt động/triển khai; ghi rõ mẫu số."],
        ["Kết quả ngoài ứng viên", "Ghi khả năng backend chặn ID ngoài danh sách. Đầu ra đã chặn không chứng minh LLM chưa từng sinh ID sai."],
        ["Độ ổn định", "Các ca đại diện gọi AI mới nhiều lần; không dùng cùng kết quả cache làm các lần lặp độc lập."],
    ])
    paragraph("Nếu chưa có nhãn chuyên gia: chỉ công bố độ khớp với nhãn kỹ thuật theo kịch bản, ghi rõ hạn chế; chưa tuyên bố độ chính xác y khoa. Không dùng cùng LLM tự gán nhãn rồi chấm chính nó như nguồn chuẩn duy nhất.")
    paragraph("Nếu LLM lỗi và chuyển fallback, giữ ca đó trong thống kê lỗi toàn pipeline và đánh dấu nguồn thực thi. Không đưa fallback vào nhóm AI thành công để tăng điểm.")

    heading("6. Mẫu bảng kết quả — chưa đo")
    table(["Kịch bản", "Số mẫu", "p50 (ms)", "p95 (ms)", "Lỗi (%)"], [
        ["AI mới theo hồ sơ", "Chưa đo", "—", "—", "—"],
        ["AI mới theo hành vi", "Chưa đo", "—", "—", "—"],
        ["Đọc cache", "Chưa đo", "—", "—", "—"],
        ["Fallback", "Chưa đo", "—", "—", "—"],
    ])
    table(["Phương pháp", "Số ca", "P@3", "Hit@3", "NDCG@3"], [
        ["Baseline theo quy tắc", "Chưa chạy", "—", "—", "—"],
        ["Xếp hạng bằng LLM", "Chưa chạy", "—", "—", "—"],
    ])
    paragraph("Dấu — nghĩa là chưa có số liệu, không phải 0. Kết quả đạt của unit test hoặc dữ liệu tracking mô phỏng không được thay vào các ô này.")

    heading("7. Điều kiện tái lập và thứ tự thực hiện")
    for text in [
        "Bước 1: xác minh backend, PostgreSQL, Redis và API key; kiểm kê số gói đang triển khai và phiên bản mã nguồn. Chỉ ghi sự hiện diện/cấu hình cần thiết, không đưa API key hoặc mật khẩu vào báo cáo.",
        "Bước 2: xây công cụ gọi đúng service/API, tách riêng không gian cache thực nghiệm; không xóa toàn bộ Redis hoặc thay dữ liệu đang dùng để vận hành.",
        "Bước 3: chạy thăm dò, sau đó đo hiệu năng và lưu từng request thành JSON/CSV; ghi CPU, RAM, hệ điều hành, môi trường local, mức đồng thời và thời điểm chạy.",
        "Bước 4: chốt tình huống và nhãn trước khi chạy đánh giá chất lượng; lưu nhãn tham chiếu, kết quả từng ca, baseline và kết quả AI riêng biệt.",
        "Bước 5: tạo bản Word kết quả cuối, biểu đồ và phần giới hạn từ file đo gốc. Không điền số liệu tự đặt trước.",
    ]:
        paragraph(text)
    paragraph("Đầu ra thống nhất: Word (.docx) là báo cáo chính; JSON/CSV là dữ liệu gốc và bằng chứng tái lập đi kèm. Nếu môi trường hoặc nhãn chưa sẵn sàng, ghi rõ phần chưa thực hiện thay vì trình bày kế hoạch như kết quả.")

    heading("8. Tệp nguồn đã đối chiếu")
    for item in [
        "src/modules/recommendations/recommendation.service.ts",
        "src/modules/recommendations/services/candidate-retrieval.service.ts",
        "src/modules/recommendations/services/ai-ranking.service.ts",
        "src/modules/recommendations/constants/recommendation.constants.ts",
        "src/modules/recommendations/recommendation.controller.ts",
        "src/infrastructure/embedding/embedding.service.ts",
        "scripts/evaluate-accuracy.mjs",
    ]:
        paragraph(item)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    assert len(reopened.tables) == 6
    print(OUTPUT)
    print(f"Verified: {len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables")


if __name__ == "__main__":
    main()
